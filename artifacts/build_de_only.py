#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import json, subprocess
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

node = r'''
const fs = require("fs");
const vm = require("vm");
const code = fs.readFileSync("/home/workdir/artifacts/qc-lab-data.js","utf8");
const ctx = { window: {} };
vm.runInNewContext(code + "\nthis.out = JSON.stringify(window.QC_CASES);", ctx);
process.stdout.write(ctx.out);
'''
CASES = json.loads(subprocess.check_output(["node", "-e", node], text=True))

def set_run_font(run, size=11, bold=False, italic=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)

def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.6)
    s.bottom_margin = Cm(1.6)
    s.left_margin = Cm(1.8)
    s.right_margin = Cm(1.8)

hp = doc.sections[0].header.paragraphs[0]
r = hp.add_run("Đại học Ngoại thương tại TP. Hồ Chí Minh  ·  2.5 Đặc tính của thông tin tài chính hữu ích")
set_run_font(r, size=9, color="5C574E")

fp = doc.sections[0].footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("ĐỀ THẢO LUẬN — chỉ tình huống và câu hỏi  ·  Trang ")
set_run_font(r, size=9, color="5C574E")
add_page_number(fp)

p = doc.add_paragraph()
r = p.add_run("2.5. ĐẶC TÍNH CỦA THÔNG TIN TÀI CHÍNH HỮU ÍCH")
set_run_font(r, size=16, bold=True, color="9A3412")

p = doc.add_paragraph()
r = p.add_run("Bộ đề 30 case thảo luận")
set_run_font(r, size=13, bold=True, color="1E3A5F")

note = (
    "Mỗi case gồm tình huống và câu hỏi thảo luận. Không ghi sẵn tên đặc tính trên đầu đề. "
    "Khi trả lời: xác định hiện tượng kinh tế, đặc tính liên quan nhất, các đặc tính dễ nhầm, "
    "và cách trình bày / ghi nhận phù hợp với Conceptual Framework 2018 (Chương 2)."
)
p = doc.add_paragraph(note)
for run in p.runs:
    set_run_font(run, size=11, italic=True)

# index
p = doc.add_paragraph()
r = p.add_run("Danh mục")
set_run_font(r, size=12, bold=True, color="1E3A5F")

tbl = doc.add_table(rows=1, cols=3)
tbl.style = "Table Grid"
for i, h in enumerate(["Case", "Công ty", "Tên tình huống"]):
    cell = tbl.rows[0].cells[i]
    cell.text = h
    shade_cell(cell, "1E3A5F")
    for para in cell.paragraphs:
        for run in para.runs:
            set_run_font(run, size=9, bold=True, color="FFFFFF")
for c in CASES:
    row = tbl.add_row().cells
    row[0].text = str(c["id"])
    row[1].text = c["company"]
    row[2].text = c["title"]
    for cell in row:
        for para in cell.paragraphs:
            for run in para.runs:
                set_run_font(run, size=10)

doc.add_page_break()

for c in CASES:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"Case {c['id']:02d}.  {c['title']}")
    set_run_font(r, size=13, bold=True, color="9A3412")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(c["company"])
    set_run_font(r, size=11, italic=True, color="1E3A5F")

    for para_text in c["scenario"].split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(para_text if para_text else "")
        set_run_font(r, size=11)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("Discussion questions")
    set_run_font(r, size=11, bold=True, color="C45C26")

    for i, q in enumerate(c["questions"], 1):
        p = doc.add_paragraph(f"{i}. {q}")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.4)
        for run in p.runs:
            set_run_font(run, size=11)

    # extra generic closer matching lecture style, without leaking the intended QC
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run(
        "Yêu cầu thêm: Chỉ ra đặc tính liên quan nhất và giải thích vì sao các đặc tính khác "
        "chỉ là hệ quả hoặc dễ gây nhầm. Nêu cách xử lý / thuyết minh phù hợp."
    )
    set_run_font(r, size=10.5, italic=True, color="5C574E")

    if c["id"] != 30:
        # subtle divider
        p = doc.add_paragraph()
        r = p.add_run("—" * 28)
        set_run_font(r, size=8, color="E4D9C5")

out = "/home/workdir/artifacts/De-30-case-dac-tinh-thong-tin-tai-chinh.docx"
doc.save(out)
print("Wrote", out)
