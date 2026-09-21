#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import json, re, pathlib

# Load cases from JS file by a light parse
js = pathlib.Path("/home/workdir/artifacts/qc-lab-data.js").read_text(encoding="utf-8")
# Convert JS object-ish to JSON-ish is hard; instead exec via node
import subprocess, tempfile, os
node = r'''
const fs = require("fs");
const vm = require("vm");
const code = fs.readFileSync("/home/workdir/artifacts/qc-lab-data.js","utf8");
const ctx = { window: {} };
vm.runInNewContext(code + "\nthis.out = JSON.stringify(window.QC_CASES);", ctx);
process.stdout.write(ctx.out);
'''
raw = subprocess.check_output(["node", "-e", node], text=True)
CASES = json.loads(raw)

def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)

def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.8)
    s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2.0)
    s.right_margin = Cm(2.0)

# Header / footer
header = doc.sections[0].header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = hp.add_run("FTU HCMC · Conceptual Framework 2018 · Ch. 2 Qualitative characteristics")
set_run_font(r, size=9, color="5C574E")

footer = doc.sections[0].footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Lab 30 case hóc búa — Đặc tính thông tin tài chính hữu ích  ·  Trang ")
set_run_font(r, size=9, color="5C574E")

def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1); run._r.append(instr); run._r.append(fldChar2)

add_page_number(fp)

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("2.5  ĐẶC TÍNH CỦA THÔNG TIN TÀI CHÍNH HỮU ÍCH")
set_run_font(r, size=18, bold=True, color="1E3A5F")

p = doc.add_paragraph()
r = p.add_run("Bộ 30 case hóc búa + hướng dẫn đáp án  ·  Nền tảng luyện tập kèm file HTML")
set_run_font(r, size=12, italic=True, color="C45C26")

intro = (
    "Bộ đề bám format slide bài giảng (tình huống tiếng Anh, câu hỏi thảo luận, ban quản trị "
    "đưa ra một lập luận dễ nghe). Nhiều case cố ý chạm hơn một đặc tính — phần khó là chỉ ra "
    "đặc tính quyết định (decisive QC), không phải liệt kê tất cả QC có thể liên quan.\n\n"
    "Khung trả lời nên dùng mỗi case:\n"
    "(1) Hiện tượng kinh tế đang được mô tả là gì?\n"
    "(2) Thông tin nào capable of making a difference (Relevance: predictive / confirmatory / materiality)?\n"
    "(3) Depiction có complete – neutral (prudence) – free from error không? Substance hay chỉ legal form?\n"
    "(4) Enhancing (comparability, verifiability, timeliness, understandability) giúp hay bị giảm?\n"
    "(5) Cost constraint có thực sự là ràng buộc hệ thống — hay chỉ là lý do để giấu thông tin?\n\n"
    "Neo lý thuyết: Conceptual Framework 2018, CF.2.1–2.43. Enhancing không thể cứu thông tin "
    "irrelevant hoặc unfaithful (CF.2.37). Faithful representation không có nghĩa chính xác tuyệt đối "
    "(CF.2.18). Prudence không phải bất đối xứng hay chọn số bi quan nhất (CF.2.16–2.17). "
    "Materiality không có ngưỡng % thống nhất (CF.2.11)."
)
p = doc.add_paragraph(intro)
for run in p.runs:
    set_run_font(run, size=11)

# Map table
doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("Bản đồ đặc tính (dùng để đối chiếu khi làm bài)")
set_run_font(r, size=13, bold=True, color="1E3A5F")

table = doc.add_table(rows=3, cols=2)
table.style = "Table Grid"
cells = [
    ("FUNDAMENTAL — Relevance",
     "Capable of making a difference in users’ decisions.\n• Predictive value (CF.2.8)\n• Confirmatory value (CF.2.9)\n• Materiality = entity-specific aspect of relevance, nature or magnitude or both (CF.2.11). Omitting, misstating or obscuring."),
    ("FUNDAMENTAL — Faithful representation",
     "Depict the substance of the phenomenon (CF.2.12).\n• Completeness (CF.2.14)\n• Neutrality, supported by prudence (CF.2.15–2.17)\n• Free from error in description and in the process — not perfect accuracy (CF.2.18–2.19)\nTrade-off với relevance: CF.2.20–2.22"),
    ("ENHANCING",
     "Comparability (≠ uniformity; consistency is a means) CF.2.24–2.29\nVerifiability (direct / indirect; range can be verified) CF.2.30–2.32\nTimeliness CF.2.33\nUnderstandability — không được loại thông tin phức tạp CF.2.34–2.36"),
    ("COST CONSTRAINT",
     "Ràng buộc pervasive khi IASB thiết kế chuẩn mực (CF.2.39–2.43).\nKhông phải thẻ miễn trừ để một entity bỏ thông tin chiếm 40% tài sản."),
    ("THỨ TỰ ÁP DỤNG",
     "CF.2.21: hiện tượng → thông tin most relevant → có faithful được không?\nCF.2.37: enhancing không biến thông tin sai/không liên quan thành useful."),
    ("BẪY THƯỜNG GẶP",
     "• % nhỏ = immaterial\n• Ước tính khác thực tế sau này = error\n• Prudence = chọn số thấp/cao nhất\n• Chưa ghi nhận = không relevant\n• User thích EBITDA = relevance thắng\n• Copy peers = comparability"),
]
for i, (h, b) in enumerate(cells):
    cell = table.rows[i // 2].cells[i % 2]
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(h)
    set_run_font(r, size=10, bold=True, color="1E3A5F")
    p2 = cell.add_paragraph()
    r = p2.add_run(b)
    set_run_font(r, size=9)
    if i % 2 == 0:
        shade(cell, "F3EADC")
    else:
        shade(cell, "FFFDF7")

doc.add_page_break()

# Index
p = doc.add_paragraph()
r = p.add_run("Mục lục 30 case")
set_run_font(r, size=14, bold=True, color="1E3A5F")

idx = doc.add_table(rows=1, cols=4)
idx.style = "Table Grid"
hdr = ["#", "Case", "Công ty", "QC quyết định"]
for i, h in enumerate(hdr):
    idx.rows[0].cells[i].text = h
    shade(idx.rows[0].cells[i], "1E3A5F")
    for para in idx.rows[0].cells[i].paragraphs:
        for run in para.runs:
            set_run_font(run, size=9, bold=True, color="FFFFFF")

for c in CASES:
    row = idx.add_row().cells
    row[0].text = str(c["id"])
    row[1].text = c["title"]
    row[2].text = c["company"]
    row[3].text = c["primary"]
    for cell in row:
        for para in cell.paragraphs:
            for run in para.runs:
                set_run_font(run, size=9)

doc.add_page_break()

for c in CASES:
    p = doc.add_paragraph()
    r = p.add_run(f"CASE {c['id']:02d}.  {c['title']}")
    set_run_font(r, size=14, bold=True, color="9A3412")

    p = doc.add_paragraph()
    r = p.add_run(f"{c['company']}   ·   {c['difficulty']}   ·   QC quyết định: {c['primary']}")
    set_run_font(r, size=10, italic=True, color="1E3A5F")

    if c.get("secondary"):
        p = doc.add_paragraph()
        r = p.add_run("Cũng chạm (nhưng thường là hệ quả / nhiễu): " + "; ".join(c["secondary"]))
        set_run_font(r, size=10, color="5C574E")

    p = doc.add_paragraph()
    r = p.add_run("Tình huống")
    set_run_font(r, size=11, bold=True, color="C45C26")

    p = doc.add_paragraph(c["scenario"])
    for run in p.runs:
        set_run_font(run, size=11)

    p = doc.add_paragraph()
    r = p.add_run("Discussion questions")
    set_run_font(r, size=11, bold=True, color="C45C26")
    for i, q in enumerate(c["questions"], 1):
        p = doc.add_paragraph(f"{i}. {q}")
        for run in p.runs:
            set_run_font(run, size=11)

    p = doc.add_paragraph()
    r = p.add_run("Gợi ý trả lời")
    set_run_font(r, size=11, bold=True, color="2F6F4E")
    p = doc.add_paragraph(c["answer"])
    for run in p.runs:
        set_run_font(run, size=10.5)

    # thin spacer
    if c["id"] != 30:
        doc.add_paragraph()

out = "/home/workdir/artifacts/30-case-dac-tinh-thong-tin-tai-chinh.docx"
doc.save(out)
print("Wrote", out, "cases", len(CASES))
